from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


SOURCE = Path('output/documents/zaria_area_command_hq_digital_reporting_case_study.docx')
OUTPUT = Path('output/pdf/zaria_area_command_hq_digital_reporting_case_study.pdf')


def iter_blocks(parent):
    for child in parent.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield 'paragraph', child
        elif isinstance(child, CT_Tbl):
            yield 'table', child


def paragraph_from_xml(document, element):
    for paragraph in document.paragraphs:
        if paragraph._p is element:
            return paragraph
    return None


def table_from_xml(document, element):
    for table in document.tables:
        if table._tbl is element:
            return table
    return None


def on_page(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor('#D9D9D9'))
    canvas.line(0.7 * inch, A4[1] - 0.53 * inch, A4[0] - 0.7 * inch, A4[1] - 0.53 * inch)
    canvas.setFillColor(colors.HexColor('#59687D'))
    canvas.setFont('Helvetica', 8)
    canvas.drawRightString(A4[0] - 0.7 * inch, A4[1] - 0.42 * inch, 'Zaria Area Command HQ Digital Reporting Case Study')
    canvas.drawCentredString(A4[0] / 2, 0.38 * inch, f'Academic prototype - not an official Nigeria Police Force service | Page {doc.page}')
    canvas.restoreState()


def build_pdf():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    source = Document(SOURCE)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle('DocTitle', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=23, leading=28, alignment=TA_CENTER, textColor=colors.black, spaceAfter=14))
    styles.add(ParagraphStyle('Subtitle', parent=styles['Normal'], fontName='Helvetica-Oblique', fontSize=12.5, leading=17, alignment=TA_CENTER, textColor=colors.HexColor('#374151'), spaceAfter=16))
    styles.add(ParagraphStyle('Body', parent=styles['Normal'], fontName='Helvetica', fontSize=9.8, leading=13.2, textColor=colors.HexColor('#1F2937'), spaceAfter=7))
    styles.add(ParagraphStyle('HeadingOne', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=15, leading=19, textColor=colors.black, spaceBefore=12, spaceAfter=7, keepWithNext=True))
    styles.add(ParagraphStyle('HeadingTwo', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11.5, leading=14, textColor=colors.black, spaceBefore=10, spaceAfter=5, keepWithNext=True))
    styles.add(ParagraphStyle('GuideBullet', parent=styles['Body'], leftIndent=17, firstLineIndent=-9, bulletIndent=7, spaceAfter=4))
    styles.add(ParagraphStyle('Cover', parent=styles['Body'], alignment=TA_CENTER, fontSize=11, leading=16, spaceAfter=10))

    story = []
    in_cover = True
    title_seen = False
    for kind, element in iter_blocks(source):
        if kind == 'paragraph':
            paragraph = paragraph_from_xml(source, element)
            if paragraph is None:
                continue
            text = paragraph.text.strip()
            style_name = paragraph.style.name
            if not text:
                story.append(Spacer(1, 5))
                continue
            if style_name == 'Title':
                story.append(Spacer(1, 1.25 * inch))
                story.append(Paragraph(escape(text), styles['DocTitle']))
                title_seen = True
                continue
            if title_seen and in_cover and style_name == 'Normal' and text.startswith('Application description'):
                story.append(Paragraph(escape(text), styles['Subtitle']))
                continue
            if style_name.startswith('Heading 1'):
                if in_cover:
                    story.append(PageBreak())
                    in_cover = False
                story.append(Paragraph(escape(text), styles['HeadingOne']))
            elif style_name.startswith('Heading 2'):
                story.append(Paragraph(escape(text), styles['HeadingTwo']))
            elif 'List Bullet' in style_name:
                story.append(Paragraph(escape(text), styles['GuideBullet'], bulletText='•'))
            else:
                target_style = styles['Cover'] if in_cover else styles['Body']
                story.append(Paragraph(escape(text).replace('\n', '<br/>'), target_style))
        else:
            table = table_from_xml(source, element)
            if table is None:
                continue
            rows = []
            for row in table.rows:
                rows.append([Paragraph(escape(cell.text), styles['Body']) for cell in row.cells])
            width = A4[0] - 1.4 * inch
            col_widths = [width / len(rows[0])] * len(rows[0])
            pdf_table = Table(rows, colWidths=col_widths, repeatRows=1, hAlign='LEFT')
            table_style = [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#17365D')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('LEADING', (0, 0), (-1, -1), 10.5),
                ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#D9D9D9')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 6),
                ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ]
            for row_index in range(1, len(rows)):
                if row_index % 2 == 0:
                    table_style.append(('BACKGROUND', (0, row_index), (-1, row_index), colors.HexColor('#F3F6FA')))
            pdf_table.setStyle(TableStyle(table_style))
            story.extend([pdf_table, Spacer(1, 8)])

    doc = SimpleDocTemplate(
        str(OUTPUT), pagesize=A4,
        leftMargin=0.7 * inch, rightMargin=0.7 * inch,
        topMargin=0.7 * inch, bottomMargin=0.62 * inch,
        title='Zaria Area Command HQ Digital Reporting Case Study',
        author='Crime Reporting System Case Study',
    )
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    print(OUTPUT.resolve())


if __name__ == '__main__':
    build_pdf()
